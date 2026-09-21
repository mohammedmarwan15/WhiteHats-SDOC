"""
Attachment readers. Every reader returns the SAME shape so the field-mapper
downstream doesn't need to know which file format it came from:

    {
        "kind": "pairs" | "text" | "unreadable",
        "pairs": [(raw_label, raw_value), ...],   # for docx/xlsx (naturally tabular)
        "text": "...",                             # for txt/pdf/ocr (line-based)
        "method": "native" | "ocr",                 # how we got the text
        "error": "..."                              # only when kind == "unreadable"
    }

Field extraction (fields.py) knows how to handle both "pairs" and "text".
"""
from pathlib import Path

# Sentinel inserted in place of PDF text we detected as scrambled (see
# _extract_pdf_text below). fields.py looks for this exact marker to tell
# "genuinely unreadable value" apart from "genuinely blank value" -- they
# need different handling downstream.
CORRUPTION_MARKER = "\ufffdCORRUPTED\ufffd"


def read_attachment(path: Path) -> dict:
    suffix = path.suffix.lower()
    try:
        if suffix == ".txt":
            return _read_txt(path)
        if suffix == ".pdf":
            return _read_pdf(path)
        if suffix == ".docx":
            return _read_docx(path)
        if suffix == ".xlsx":
            return _read_xlsx(path)
        return {"kind": "unreadable", "error": f"unsupported file type: {suffix}"}
    except Exception as exc:  # noqa: BLE001 -- deliberately broad: any failure
        # here must become a NEEDS_REVIEW case, not a crash of the whole run.
        return {"kind": "unreadable", "error": f"{type(exc).__name__}: {exc}"}


def _read_txt(path: Path) -> dict:
    data = path.read_bytes()
    if len(data) == 0:
        return {"kind": "unreadable", "error": "empty file (0 bytes)"}
    text = data.decode("utf-8", errors="replace")
    if not text.strip():
        return {"kind": "unreadable", "error": "file has no readable content"}
    return {"kind": "text", "text": text, "method": "native"}


def _read_pdf(path: Path) -> dict:
    import pdfplumber

    if path.stat().st_size == 0:
        return {"kind": "unreadable", "error": "empty file (0 bytes)"}

    try:
        with pdfplumber.open(path) as pdf:
            text = _extract_pdf_text(pdf)
    except Exception as exc:
        return {"kind": "unreadable", "error": f"garbled/corrupt PDF: {exc}"}

    if text.strip():
        return {"kind": "text", "text": text, "method": "native"}

    # No text layer at all -> likely an image-only scan. Fall back to OCR.
    return _ocr_pdf(path)


def _extract_pdf_text(pdf, overlap_threshold: float = -1.0, line_tolerance: float = 2.0) -> str:
    """Rebuilds each page's text from individual character positions instead
    of calling pdfplumber's extract_text() directly, so we can catch a real
    rendering defect found in a handful of this dataset's PDFs: two separate
    pieces of text (e.g. a label's tail end and the field's actual value)
    drawn on top of each other at overlapping x-positions, which produces
    scrambled text no amount of regex/alias matching can parse correctly --
    e.g. "Notify Party/Intermediate ConsKiTgPne CeO., LTD" is "Consignee"
    and "KTP CO., LTD" interleaved character-by-character.

    In a cleanly-rendered PDF, consecutive characters on the same line never
    overlap by more than a fraction of a point (touching/kerning gives ~0).
    A gap more negative than `overlap_threshold` means two glyphs are
    genuinely stacked on each other -- a corruption signature, not a
    formatting quirk. When we see that, we cut the line there and drop in
    CORRUPTION_MARKER instead of guessing at scrambled text.
    """
    lines_out = []
    for page in pdf.pages:
        chars = sorted(page.chars, key=lambda c: (c["top"], c["x0"]))
        grouped, cur_line, cur_top = [], [], None
        for c in chars:
            if cur_top is None or abs(c["top"] - cur_top) <= line_tolerance:
                cur_line.append(c)
                cur_top = c["top"] if cur_top is None else cur_top
            else:
                grouped.append(cur_line)
                cur_line, cur_top = [c], c["top"]
        if cur_line:
            grouped.append(cur_line)

        for line_chars in grouped:
            line_chars.sort(key=lambda c: c["x0"])
            kept, prev_x1, corrupted = [], None, False
            for c in line_chars:
                if prev_x1 is not None and (c["x0"] - prev_x1) < overlap_threshold:
                    corrupted = True
                    break
                kept.append(c["text"])
                prev_x1 = c["x1"]
            line_text = "".join(kept)
            if corrupted:
                line_text = (line_text.rstrip() + " " + CORRUPTION_MARKER).strip()
            lines_out.append(line_text)
    return "\n".join(lines_out)


def _ocr_pdf(path: Path) -> dict:
    try:
        import pypdfium2 as pdfium
        import pytesseract

        pdf = pdfium.PdfDocument(str(path))
        texts = []
        for i in range(len(pdf)):
            page = pdf[i]
            bitmap = page.render(scale=2.0)
            pil_image = bitmap.to_pil()
            texts.append(pytesseract.image_to_string(pil_image))
        text = "\n".join(texts)
    except Exception as exc:
        return {"kind": "unreadable", "error": f"OCR failed: {exc}"}

    if not text.strip():
        return {"kind": "unreadable", "error": "image-only PDF, OCR found no text"}
    return {"kind": "text", "text": text, "method": "ocr"}


def _read_docx(path: Path) -> dict:
    import docx

    d = docx.Document(str(path))
    pairs = []
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0]:
                pairs.append((cells[0], cells[1]))
    # also keep plain paragraphs as backup text (e.g. B/L NO. line lives
    # outside any table in this template)
    para_text = "\n".join(p.text for p in d.paragraphs if p.text.strip())

    if not pairs and not para_text.strip():
        return {"kind": "unreadable", "error": "docx has no readable content"}
    return {"kind": "pairs", "pairs": pairs, "text": para_text, "method": "native"}


def _read_xlsx(path: Path) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    pairs = []
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        cells = [str(c).strip() if c is not None else "" for c in row]
        if len(cells) >= 2 and cells[0] and cells[1]:
            pairs.append((cells[0], cells[1]))
    if not pairs:
        return {"kind": "unreadable", "error": "xlsx has no readable label/value rows"}
    return {"kind": "pairs", "pairs": pairs, "method": "native"}